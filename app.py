from flask import Flask, render_template, request, redirect, url_for, session, flash
import sqlite3
import pandas as pd
import numpy as np
import pickle
from werkzeug.security import generate_password_hash, check_password_hash
import os
from io import BytesIO
from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

app = Flask(__name__)
app.secret_key = "secret123"

# Load model and scaler
model = pickle.load(open("model.pkl", "rb"))
scaler = pickle.load(open("scaler.pkl", "rb"))

# SHAP explainer
import shap
explainer = shap.TreeExplainer(model)

FEATURE_NAMES = ["Knowledge Test", "Study Hours", "Previous Marks", "Assignments"]


# ---------- HELPER: SAFE FLOAT ----------
def safe_float(v):
    try:
        if v is None or str(v).strip() == "":
            return None
        return float(v)
    except (ValueError, TypeError):
        return None


# ---------- HELPER: VALIDATE + NORMALIZE ----------
def validate_and_normalize(kt, sh, pm, asg, marks_range="100"):
    if sh is not None and not (0 <= sh <= 24):
        return None, "Study Hours must be between 0 and 24."

    max_mark = 50 if marks_range == "50" else 100
    for label, val in [("Knowledge Test", kt),
                       ("Previous Marks", pm),
                       ("Assignments", asg)]:
        if val is not None and not (0 <= val <= max_mark):
            return None, f"{label} must be between 0 and {max_mark}."

    scale = 100 / max_mark if max_mark == 50 else 1.0
    return {
        "kt":  round(kt * scale, 2)  if kt  is not None else None,
        "sh":  sh,
        "pm":  round(pm * scale, 2)  if pm  is not None else None,
        "asg": round(asg * scale, 2) if asg is not None else None,
    }, None


# ---------- HELPER: AUTO-DETECT RANGE ----------
def detect_marks_range(kt, pm, asg):
    values = [v for v in (kt, pm, asg) if v is not None]
    if not values:
        return "100"
    return "50" if max(values) <= 50 else "100"


# ---------- HELPER: PREDICT + CONFIDENCE ----------
def predict_student(kt, sh, pm, asg):
    """Returns (label, confidence_percent). N/A with 0% if input missing."""
    if None in (kt, sh, pm, asg):
        return "N/A", 0
    features = np.array([[kt, sh, pm, asg]])
    features_scaled = scaler.transform(features)
    pred = model.predict(features_scaled)[0]
    proba = model.predict_proba(features_scaled)[0]
    confidence = round(float(max(proba)) * 100, 1)
    label = "PASS" if pred == 1 else "FAIL"
    return label, confidence


# ---------- HELPER: SHAP ----------
def explain_student(kt, sh, pm, asg):
    if None in (kt, sh, pm, asg):
        return None

    features = np.array([[kt, sh, pm, asg]])
    features_scaled = scaler.transform(features)
    shap_values = explainer.shap_values(features_scaled)

    if isinstance(shap_values, list):
        contribs = shap_values[1][0]
    else:
        contribs = shap_values[0, :, 1]

    shap_list = []
    for i, name in enumerate(FEATURE_NAMES):
        val = float(contribs[i])
        shap_list.append({
            "name": name,
            "value": round(val, 3),
            "direction": "PASS" if val > 0 else "FAIL",
            "abs": abs(val)
        })
    shap_list.sort(key=lambda x: x["abs"], reverse=True)
    return shap_list


# ---------------- DATABASE ----------------
def init_db():
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE,
        email TEXT,
        password TEXT
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS predictions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT,
        knowledge_test REAL,
        study_hours REAL,
        previous_marks REAL,
        assignments REAL,
        prediction TEXT,
        confidence REAL,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id TEXT UNIQUE,
        name TEXT NOT NULL,
        knowledge_test REAL,
        study_hours REAL,
        previous_marks REAL,
        assignments REAL,
        prediction TEXT,
        confidence REAL,
        added_by TEXT,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
    )""")
    conn.commit()
    conn.close()


init_db()


# ---------------- ROUTES ----------------

@app.route("/")
def index():
    return redirect(url_for("login"))


# ---------- REGISTER ----------
@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form["username"]
        email = request.form["email"]
        password = generate_password_hash(request.form["password"])
        try:
            conn = sqlite3.connect("users.db")
            c = conn.cursor()
            c.execute("INSERT INTO users (username, email, password) VALUES (?, ?, ?)",
                      (username, email, password))
            conn.commit()
            conn.close()
            flash("Registration successful! Please login.", "success")
            return redirect(url_for("login"))
        except sqlite3.IntegrityError:
            flash("Username already exists. Try another.", "danger")
    return render_template("register.html")


# ---------- LOGIN ----------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]
        conn = sqlite3.connect("users.db")
        c = conn.cursor()
        c.execute("SELECT * FROM users WHERE username = ?", (username,))
        user = c.fetchone()
        conn.close()
        if user and check_password_hash(user[3], password):
            session["username"] = username
            flash("Login successful!", "success")
            return redirect(url_for("home"))
        else:
            flash("Invalid username or password.", "danger")
    return render_template("login.html")


# ---------- HOME ----------
@app.route("/home")
def home():
    if "username" not in session:
        return redirect(url_for("login"))
    return render_template("home.html", username=session["username"])


# ---------- PREDICTION ----------
@app.route("/prediction", methods=["GET", "POST"])
def prediction():
    if "username" not in session:
        return redirect(url_for("login"))

    result = None
    error = None

    if request.method == "POST":
        marks_range = request.form.get("marks_range", "100")

        kt  = safe_float(request.form.get("knowledge_test"))
        sh  = safe_float(request.form.get("study_hours"))
        pm  = safe_float(request.form.get("previous_marks"))
        asg = safe_float(request.form.get("assignments"))

        normalized, err = validate_and_normalize(kt, sh, pm, asg, marks_range)
        if err:
            error = err
        else:
            kt_n = normalized["kt"]
            sh_n = normalized["sh"]
            pm_n = normalized["pm"]
            asg_n = normalized["asg"]

            pred_label, confidence = predict_student(kt_n, sh_n, pm_n, asg_n)
            shap_list = explain_student(kt_n, sh_n, pm_n, asg_n)

            conn = sqlite3.connect("users.db")
            c = conn.cursor()
            c.execute("""INSERT INTO predictions
                         (username, knowledge_test, study_hours,
                          previous_marks, assignments, prediction, confidence)
                         VALUES (?, ?, ?, ?, ?, ?, ?)""",
                      (session["username"], kt_n, sh_n, pm_n, asg_n, pred_label, confidence))
            conn.commit()
            conn.close()

            tips = []
            if kt_n < 70:
                tips.append("Improve knowledge test performance through regular practice.")
            if sh_n < 3:
                tips.append("Increase daily study hours to at least 3-4 hours.")
            if pm_n < 60:
                tips.append("Revise previous academic concepts.")
            if asg_n < 75:
                tips.append("Complete assignments regularly and on time.")
            if pred_label == "PASS":
                tips.append("Great! Maintain your current study habits.")
            else:
                tips.append("Focus on weak areas and seek academic support.")

            result = {
                "prediction": pred_label,
                "confidence": confidence,
                "recommendation": tips,
                "input": {
                    "kt": kt, "sh": sh, "pm": pm, "asg": asg,
                    "kt_n": kt_n, "sh_n": sh_n, "pm_n": pm_n, "asg_n": asg_n,
                },
                "shap": shap_list,
                "marks_range": marks_range,
            }

    return render_template("prediction.html", result=result, error=error)


# ---------- ABOUT ----------
@app.route("/about")
def about():
    return render_template("about.html")


# ---------- LOGOUT ----------
@app.route("/logout")
def logout():
    session.pop("username", None)
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))


# ---------- BULK ----------
@app.route("/bulk")
def bulk():
    if "username" not in session:
        return redirect(url_for("login"))
    return render_template("bulk.html")


@app.route("/download_template")
def download_template():
    if "username" not in session:
        return redirect(url_for("login"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = ["Student ID", "Name", "Knowledge Test", "Study Hours", "Previous Marks", "Assignments"]
    ws.append(headers)
    ws.append(["S001", "Sample Student 1", 85, 3, 72, 90])
    ws.append(["S002", "Sample Student 2", 60, 1, 48, 55])

    for col in range(1, 7):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")

    widths = [12, 22, 16, 14, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="student_template.xlsx"
    )


@app.route("/bulk_predict", methods=["POST"])
def bulk_predict():
    if "username" not in session:
        return redirect(url_for("login"))

    if "file" not in request.files:
        flash("No file uploaded.", "danger")
        return redirect(url_for("bulk"))

    file = request.files["file"]
    if not file.filename.endswith(".xlsx"):
        flash("Please upload an .xlsx file.", "danger")
        return redirect(url_for("bulk"))

    try:
        df = pd.read_excel(file)

        required = ["Knowledge Test", "Study Hours", "Previous Marks", "Assignments"]
        for col in required:
            if col not in df.columns:
                flash(f"Missing column: {col}. Please use the template.", "danger")
                return redirect(url_for("bulk"))

        predictions = []
        confidences = []
        recommendations = []

        for _, row in df.iterrows():
            features = np.array([[
                float(row["Knowledge Test"]),
                float(row["Study Hours"]),
                float(row["Previous Marks"]),
                float(row["Assignments"])
            ]])
            features_scaled = scaler.transform(features)
            pred = model.predict(features_scaled)[0]
            proba = model.predict_proba(features_scaled)[0]
            conf = round(float(max(proba)) * 100, 1)

            pred_label = "PASS" if pred == 1 else "FAIL"
            predictions.append(pred_label)
            confidences.append(conf)

            tips = []
            if row["Knowledge Test"] < 70:
                tips.append("Improve test score")
            if row["Study Hours"] < 3:
                tips.append("Study more")
            if row["Previous Marks"] < 60:
                tips.append("Revise basics")
            if row["Assignments"] < 75:
                tips.append("Complete assignments")
            if not tips:
                tips.append("Maintain current habits")
            recommendations.append("; ".join(tips))

        df["Prediction"] = predictions
        df["Confidence %"] = confidences
        df["Recommendation"] = recommendations

        wb = Workbook()
        ws = wb.active
        ws.title = "Results"

        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

        for col in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")

        pred_col = list(df.columns).index("Prediction") + 1
        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=pred_col)
            if cell.value == "PASS":
                cell.font = Font(bold=True, color="28A745")
            else:
                cell.font = Font(bold=True, color="DC3545")

        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 3

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="prediction_results.xlsx"
        )

    except Exception as e:
        flash(f"Error processing file: {str(e)}", "danger")
        return redirect(url_for("bulk"))


# ---------- STUDENTS ----------
@app.route("/students")
def students():
    if "username" not in session:
        return redirect(url_for("login"))

    q = request.args.get("q", "").strip()
    show_all = request.args.get("show_all", "") == "1"

    rows = []
    if q or show_all:
        conn = sqlite3.connect("users.db")
        c = conn.cursor()
        if q:
            c.execute("""SELECT id, student_id, name, knowledge_test, study_hours,
                                previous_marks, assignments, prediction, confidence, timestamp
                         FROM students
                         WHERE name LIKE ? OR student_id LIKE ?
                         ORDER BY name""", (f"%{q}%", f"%{q}%"))
        else:
            c.execute("""SELECT id, student_id, name, knowledge_test, study_hours,
                                previous_marks, assignments, prediction, confidence, timestamp
                         FROM students ORDER BY name""")
        rows = c.fetchall()
        conn.close()

    return render_template("students.html", students=rows, q=q, show_all=show_all)


# ---------- ADD STUDENT ----------
@app.route("/students/add", methods=["GET", "POST"])
def add_student():
    if "username" not in session:
        return redirect(url_for("login"))

    error = None
    if request.method == "POST":
        sid = request.form.get("student_id", "").strip() or None
        name = request.form["name"].strip()
        marks_range = request.form.get("marks_range", "100")

        kt  = safe_float(request.form.get("knowledge_test"))
        sh  = safe_float(request.form.get("study_hours"))
        pm  = safe_float(request.form.get("previous_marks"))
        asg = safe_float(request.form.get("assignments"))

        normalized, err = validate_and_normalize(kt, sh, pm, asg, marks_range)
        if err:
            error = err
        else:
            kt_n = normalized["kt"]
            sh_n = normalized["sh"]
            pm_n = normalized["pm"]
            asg_n = normalized["asg"]

            pred_label, confidence = predict_student(kt_n, sh_n, pm_n, asg_n)

            try:
                conn = sqlite3.connect("users.db")
                c = conn.cursor()
                c.execute("""INSERT INTO students
                    (student_id, name, knowledge_test, study_hours,
                     previous_marks, assignments, prediction, confidence, added_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (sid, name, kt_n, sh_n, pm_n, asg_n, pred_label, confidence, session["username"]))
                conn.commit()
                conn.close()
                flash(f"Student '{name}' added successfully!", "success")
                return redirect(url_for("students"))
            except sqlite3.IntegrityError:
                flash("Student ID already exists.", "danger")

    return render_template("add_student.html", error=error)


# ---------- VIEW STUDENT ----------
@app.route("/students/<int:sid>")
def student_detail(sid):
    if "username" not in session:
        return redirect(url_for("login"))
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("SELECT * FROM students WHERE id = ?", (sid,))
    student = c.fetchone()
    conn.close()
    if not student:
        flash("Student not found.", "danger")
        return redirect(url_for("students"))

    shap_list = explain_student(student[3], student[4], student[5], student[6])
    return render_template("student_detail.html", student=student, shap=shap_list)


# ---------- DELETE STUDENT ----------
@app.route("/students/delete/<int:sid>")
def delete_student(sid):
    if "username" not in session:
        return redirect(url_for("login"))
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("DELETE FROM students WHERE id = ?", (sid,))
    conn.commit()
    conn.close()
    flash("Student deleted.", "success")
    return redirect(url_for("students"))


# ---------- IMPORT ----------
@app.route("/students/import", methods=["POST"])
def import_students():
    if "username" not in session:
        return redirect(url_for("login"))

    file = request.files.get("file")
    if not file or not file.filename.endswith(".xlsx"):
        flash("Please upload a valid .xlsx file.", "danger")
        return redirect(url_for("students"))

    summary = {
        "total": 0, "added": 0, "skipped_duplicate": 0, "skipped_no_name": 0,
        "skipped_out_of_range": 0, "pass_count": 0, "fail_count": 0,
        "na_count": 0, "error": None, "detected_range": None,
    }

    try:
        df = pd.read_excel(file)
        required = ["Name", "Knowledge Test", "Study Hours", "Previous Marks", "Assignments"]
        for col in required:
            if col not in df.columns:
                summary["error"] = f"Missing column: {col}"
                return render_template("import_result.html", summary=summary)

        summary["total"] = len(df)

        all_kt  = [safe_float(v) for v in df.get("Knowledge Test", [])]
        all_pm  = [safe_float(v) for v in df.get("Previous Marks", [])]
        all_asg = [safe_float(v) for v in df.get("Assignments", [])]

        detected_max = max(
            [v for v in (all_kt + all_pm + all_asg) if v is not None] or [0]
        )
        marks_range = "50" if detected_max <= 50 else "100"
        summary["detected_range"] = marks_range

        conn = sqlite3.connect("users.db")
        c = conn.cursor()

        for _, row in df.iterrows():
            sid = str(row.get("Student ID", "")).strip() or None
            raw_name = row.get("Name")
            name = str(raw_name).strip() if pd.notna(raw_name) else ""

            if not name or name.lower() in ("nan", "none", "(dropped out)"):
                summary["skipped_no_name"] += 1
                continue

            kt = safe_float(row.get("Knowledge Test"))
            sh = safe_float(row.get("Study Hours"))
            pm = safe_float(row.get("Previous Marks"))
            asg = safe_float(row.get("Assignments"))

            normalized, err = validate_and_normalize(kt, sh, pm, asg, marks_range)
            if err:
                summary["skipped_out_of_range"] += 1
                continue

            kt_n = normalized["kt"]
            sh_n = normalized["sh"]
            pm_n = normalized["pm"]
            asg_n = normalized["asg"]

            pred_label, confidence = predict_student(kt_n, sh_n, pm_n, asg_n)

            if pred_label == "PASS":
                summary["pass_count"] += 1
            elif pred_label == "FAIL":
                summary["fail_count"] += 1
            else:
                summary["na_count"] += 1

            try:
                c.execute("""INSERT INTO students
                    (student_id, name, knowledge_test, study_hours, previous_marks,
                     assignments, prediction, confidence, added_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (sid, name, kt_n, sh_n, pm_n, asg_n, pred_label, confidence, session["username"]))
                summary["added"] += 1
            except sqlite3.IntegrityError:
                summary["skipped_duplicate"] += 1

        conn.commit()
        conn.close()

    except Exception as e:
        summary["error"] = str(e)

    return render_template("import_result.html", summary=summary)

# ---------- WHAT-IF ANALYSIS ----------
@app.route("/whatif", methods=["GET", "POST"])
def whatif():
    if "username" not in session:
        return redirect(url_for("login"))

    result = None

    if request.method == "POST":
        kt  = float(request.form.get("knowledge_test", 50))
        sh  = float(request.form.get("study_hours", 3))
        pm  = float(request.form.get("previous_marks", 50))
        asg = float(request.form.get("assignments", 50))

        pred_label, confidence = predict_student(kt, sh, pm, asg)

        result = {
            "prediction": pred_label,
            "confidence": confidence,
            "kt": kt, "sh": sh, "pm": pm, "asg": asg,
        }

    return render_template("whatif.html", result=result)

# ---------- EXPORT CLASS TO EXCEL ----------
@app.route("/students/export")
def export_students():
    if "username" not in session:
        return redirect(url_for("login"))

    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("""SELECT student_id, name, knowledge_test, study_hours,
                        previous_marks, assignments, prediction, confidence
                 FROM students ORDER BY name""")
    rows = c.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Class Report"

    headers = ["Student ID", "Name", "KT", "SH", "PM", "ASG",
               "Prediction", "Confidence %", "Top Reason 1", "Top Reason 2"]
    ws.append(headers)

    for r in rows:
        sid, name, kt, sh, pm, asg, pred, conf = r

        # Get SHAP for top 2 reasons
        shap_list = explain_student(kt, sh, pm, asg)
        reason1 = reason2 = "—"
        if shap_list and len(shap_list) >= 1:
            reason1 = f"{shap_list[0]['name']} ({shap_list[0]['value']:+})"
        if shap_list and len(shap_list) >= 2:
            reason2 = f"{shap_list[1]['name']} ({shap_list[1]['value']:+})"

        ws.append([sid, name, kt, sh, pm, asg, pred,
                   conf if conf else "—", reason1, reason2])

    # Style headers
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")

    # Color-code PASS / FAIL
    pred_col = headers.index("Prediction") + 1
    for row in range(2, len(rows) + 2):
        cell = ws.cell(row=row, column=pred_col)
        if cell.value == "PASS":
            cell.font = Font(bold=True, color="28A745")
        elif cell.value == "FAIL":
            cell.font = Font(bold=True, color="DC3545")

    # Auto-size columns
    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="class_report.xlsx"
    )

# ---------- CLASS DASHBOARD ----------
@app.route("/dashboard")
def dashboard():
    if "username" not in session:
        return redirect(url_for("login"))

    conn = sqlite3.connect("users.db")
    c = conn.cursor()

    c.execute("SELECT COUNT(*) FROM students")
    total = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM students WHERE prediction='PASS'")
    passed = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM students WHERE prediction='FAIL'")
    failed = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM students WHERE prediction='N/A'")
    na = c.fetchone()[0]

    c.execute("""SELECT AVG(knowledge_test), AVG(study_hours),
                        AVG(previous_marks), AVG(assignments) FROM students""")
    avg_row = c.fetchone()

    c.execute("SELECT AVG(confidence) FROM students WHERE confidence IS NOT NULL")
    avg_conf = c.fetchone()[0]

    conn.close()

    pass_rate = round((passed / total * 100), 1) if total > 0 else 0

    stats = {
        "total": total,
        "passed": passed,
        "failed": failed,
        "na": na,
        "pass_rate": pass_rate,
        "avg_kt": round(avg_row[0] or 0, 1),
        "avg_sh": round(avg_row[1] or 0, 1),
        "avg_pm": round(avg_row[2] or 0, 1),
        "avg_asg": round(avg_row[3] or 0, 1),
        "avg_conf": round(avg_conf or 0, 1),
    }
    return render_template("dashboard.html", stats=stats)

# ---------- EDIT STUDENT ----------
@app.route("/students/edit/<int:sid>", methods=["GET", "POST"])
def edit_student(sid):
    if "username" not in session:
        return redirect(url_for("login"))

    conn = sqlite3.connect("users.db")
    c = conn.cursor()

    if request.method == "POST":
        name = request.form["name"].strip()
        kt   = safe_float(request.form.get("knowledge_test"))
        sh   = safe_float(request.form.get("study_hours"))
        pm   = safe_float(request.form.get("previous_marks"))
        asg  = safe_float(request.form.get("assignments"))

        pred_label, confidence = predict_student(kt, sh, pm, asg)

        c.execute("""UPDATE students
                     SET name=?, knowledge_test=?, study_hours=?,
                         previous_marks=?, assignments=?,
                         prediction=?, confidence=?
                     WHERE id=?""",
                  (name, kt, sh, pm, asg, pred_label, confidence, sid))
        conn.commit()
        conn.close()
        flash(f"Student '{name}' updated.", "success")
        return redirect(url_for("students"))

    c.execute("SELECT * FROM students WHERE id=?", (sid,))
    student = c.fetchone()
    conn.close()

    if not student:
        flash("Student not found.", "danger")
        return redirect(url_for("students"))

    return render_template("edit_student.html", student=student)
# ---------------- RUN ----------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)